#!/usr/bin/env python3
"""Offline ELS workbook -> model_v31 assets. Never connects to a server."""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import shutil
import sys
import tempfile
from zipfile import BadZipFile
from dataclasses import fields, is_dataclass
from pathlib import Path

if not (3, 11) <= sys.version_info[:2] < (3, 14):
    raise SystemExit('이 도구는 Python 3.11~3.13을 사용하세요. 검증 환경은 Python 3.12입니다.')

import openpyxl
from xsdata.formats.dataclass.parsers import XmlParser
from xsdata.formats.dataclass.serializers import XmlSerializer
from xsdata.formats.dataclass.serializers.config import SerializerConfig

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from src.entities import model_v31 as m
from schema import SOURCE, SUPPLEMENT, FILES, ROW_LAYOUT

HEADERS = ['테이블', '레코드', '컬럼', '입력값', '필수', '설명']
PATH_HEADERS = ['가공레코드', '순서', 'X_mm', 'Y_mm', 'Z_mm', 'I', 'J', 'K']
NOTES = [
    '등록 구조 시험용입니다. NC 실행, 충돌 검증, ISO 전체 적합성 검증은 수행하지 않습니다.',
    'FreeformOperation + ToolpathFeature만 지원합니다. 다른 피처·전략은 생성하지 않습니다.',
    '공작물 좌표계의 경로점과 공구축은 입력값만 사용합니다. 기계축 역기구학·보간은 계산하지 않습니다.',
    '장비 축·주축·매거진의 상세 능력, 셋업 고정구, 공차·가공 전략은 이번 프로필에 포함되지 않습니다.',
    '공구 numerical_value 이름은 현재 DT 구현의 로컬 속성명이며 ISO 13399 사전 코드 인증을 뜻하지 않습니다.',
]


class InputError(Exception):
    pass


def parse_value(value, kind):
    if kind == 'text':
        if not isinstance(value, str):
            raise ValueError('문자열로 입력해야 합니다. ID의 앞자리 0을 보존하세요.')
        if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value):
            raise ValueError('XML에 사용할 수 없는 제어문자가 있습니다.')
        return value.strip()
    if kind == 'bool':
        if isinstance(value, bool):
            return value
        if isinstance(value, str) and value.strip().lower() in ('true', 'false'):
            return value.strip().lower() == 'true'
        raise ValueError('true 또는 false로 입력하세요.')
    if kind in ('vector', 'unitvector'):
        if not isinstance(value, str):
            raise ValueError('쉼표로 구분한 숫자 3개를 입력하세요.')
        result = [float(x.strip()) for x in value.split(',')]
        if len(result) != 3 or not all(math.isfinite(x) for x in result):
            raise ValueError('유한한 숫자 3개가 필요합니다.')
        if kind == 'unitvector' and not math.isclose(sum(x*x for x in result), 1, abs_tol=1e-6):
            raise ValueError('길이가 1인 방향벡터가 필요합니다.')
        return result
    if kind in ('number', 'positive', 'nonnegative', 'integer'):
        if isinstance(value, bool):
            raise ValueError('숫자가 필요합니다.')
        result = float(value)
        if not math.isfinite(result):
            raise ValueError('유한한 숫자가 필요합니다.')
        if kind in ('positive', 'integer') and result <= 0:
            raise ValueError('0보다 큰 값이 필요합니다.')
        if kind == 'nonnegative' and result < 0:
            raise ValueError('0 이상이어야 합니다.')
        if kind == 'integer':
            if not result.is_integer():
                raise ValueError('정수가 필요합니다.')
            return int(result)
        return result
    if not isinstance(value, str) or value.strip() not in kind.split('|'):
        raise ValueError('허용값: ' + kind)
    return value.strip()


class Input:
    def __init__(self):
        self.data = {}
        self.locations = {}
        self.errors = []
        self.checks = []
        self.unmapped = []
        self.paths = {}

    def error(self, table, record, column, message):
        location = self.locations.get((table, record, column), '필수보완 또는 업체입력')
        self.errors.append({'cell': location, 'table': table, 'record': record,
                            'column': column, 'message': message})

    def read(self, filename):
        # Keep previously delivered vertical workbooks readable, while new
        # workbooks use database-style rows and explicit foreign keys.
        probe = openpyxl.load_workbook(filename, data_only=False, read_only=True)
        row_format = 'Part_Info' in probe.sheetnames
        original_format = row_format and (probe['Part_Info']['A1'].value == 'Part_No' or 'DT_공정' in probe.sheetnames)
        probe.close()
        if original_format:
            from original_reader import read_original
            return read_original(self, filename, parse_value)
        if row_format:
            return self.read_rows(filename)
        wb = openpyxl.load_workbook(filename, data_only=False, read_only=True)
        try:
            for sheet, schema in [('업체입력', SOURCE), ('필수보완', SUPPLEMENT), ('파일입력', {'Files': FILES})]:
                if sheet not in wb.sheetnames:
                    self.errors.append({'cell': sheet, 'message': '시트가 없습니다.'})
                    continue
                ws = wb[sheet]
                if [c.value for c in ws[4]][:6] != HEADERS:
                    self.errors.append({'cell': sheet + '!A4:F4', 'message': '양식의 헤더가 변경되었습니다.'})
                    continue
                seen = set()
                for cells in ws.iter_rows(min_row=5, max_col=6):
                    table, record, column, value = [c.value for c in cells[:4]]
                    if not any(v is not None for v in (table, record, column, value)):
                        continue
                    location = f'{sheet}!D{cells[0].row}'
                    if any(c.data_type == 'f' for c in cells[:4]):
                        self.errors.append({'cell': location, 'message': '수식은 지원하지 않습니다. 값으로 붙여넣으세요.'})
                        continue
                    if not all(isinstance(x, str) and x.strip() for x in (table, record, column)):
                        self.errors.append({'cell': location, 'message': '테이블·레코드·컬럼을 문자열로 입력하세요.'})
                        continue
                    table, record, column = table.strip(), record.strip(), column.strip()
                    key = table, record, column
                    self.locations[key] = location
                    definitions = {row[0]: row for row in schema.get(table, [])}
                    if column not in definitions:
                        self.unmapped.append({'cell': location, 'table': table, 'record': record, 'column': column,
                                              'reason': '이 시험 프로필에서 사용하지 않는 컬럼'})
                        continue
                    if key in seen:
                        self.error(*key, '같은 레코드의 컬럼이 중복되었습니다.')
                        continue
                    seen.add(key)
                    row = self.data.setdefault(table, {}).setdefault(record, {})
                    if value is None or (isinstance(value, str) and not value.strip()):
                        continue
                    try:
                        row[column] = parse_value(value, definitions[column][1])
                        self.checks.append({'cell': location, 'table': table, 'record': record, 'column': column,
                                            'required': definitions[column][2], 'mapping': definitions[column][3]})
                    except (ValueError, TypeError, OverflowError):
                        self.error(*key, '입력 형식/범위 오류. 필요한 형식: ' + definitions[column][1])
            for table, definitions in {**SOURCE, 'Files': FILES}.items():
                rows = self.data.get(table, {})
                if not rows and table != 'Files':
                    self.errors.append({'cell': '업체입력', 'table': table, 'message': '최소 1개 레코드가 필요합니다.'})
                for record, row in rows.items():
                    for col, _, required, _, _ in definitions + SUPPLEMENT.get(table, []):
                        if required and col not in row:
                            self.error(table, record, col, '필수값이 없습니다.')
            # Supplement-only records must not silently create a new source record.
            for table in SUPPLEMENT:
                for record in self.data.get(table, {}):
                    if not any(c['table'] == table and c['record'] == record and c['cell'].startswith('업체입력!') for c in self.checks):
                        self.error(table, record, '', '업체입력에 대응하는 레코드가 없습니다.')
            if '공구경로' in wb.sheetnames:
                ws = wb['공구경로']
                if [c.value for c in ws[4]][:8] != PATH_HEADERS:
                    self.errors.append({'cell': '공구경로!A4:H4', 'message': '경로 헤더가 변경되었습니다.'})
                else:
                    for row_number, cells in enumerate(ws.iter_rows(min_row=5, max_col=8), 5):
                        values = [c.value for c in cells]
                        if all(v is None for v in values):
                            continue
                        try:
                            if any(c.data_type == 'f' for c in cells):
                                raise ValueError()
                            rec = parse_value(values[0], 'text')
                            seq = parse_value(values[1], 'integer')
                            xyzijk = [parse_value(v, 'number') for v in values[2:]]
                            if not math.isclose(sum(x*x for x in xyzijk[3:]), 1, abs_tol=1e-6):
                                raise ValueError()
                            self.paths.setdefault(rec, []).append((seq, xyzijk, f'공구경로!A{row_number}'))
                        except (ValueError, TypeError, OverflowError):
                            self.errors.append({'cell': f'공구경로!A{row_number}:H{row_number}', 'message': '레코드/순서/좌표/단위 공구축을 확인하세요. 수식은 불가합니다.'})
        finally:
            wb.close()

    def read_rows(self, filename):
        wb = openpyxl.load_workbook(filename, data_only=False, read_only=True)
        source_records = set()
        supplements = []
        try:
            for sheet, (table, definitions, supplement) in ROW_LAYOUT.items():
                if sheet not in wb.sheetnames:
                    self.errors.append({'cell': sheet, 'message': '필요한 시트가 없습니다.'})
                    continue
                ws = wb[sheet]
                headers = [cell.value for cell in ws[4]]
                nonempty = [h for h in headers if h is not None]
                if any(not isinstance(h, str) for h in nonempty) or len(nonempty) != len(set(nonempty)):
                    self.errors.append({'cell': sheet+'!4', 'message': '컬럼 헤더가 중복되었거나 문자열이 아닙니다.'})
                    continue
                expected = ['Record_ID'] + [d[0] for d in definitions]
                if any(h not in headers for h in expected):
                    self.errors.append({'cell': sheet+'!4', 'message': '필요한 컬럼 헤더가 없습니다. 양식 헤더를 유지하세요.'})
                    continue
                id_index = headers.index('Record_ID')
                defs = {d[0]: d for d in definitions}
                seen = set()
                for row_number, cells in enumerate(ws.iter_rows(min_row=5), 5):
                    if all(c.value is None or c.value == '' for c in cells):
                        continue
                    id_cell = cells[id_index]
                    record = id_cell.value
                    if id_cell.data_type == 'f' or not isinstance(record, str) or not record.strip():
                        self.errors.append({'cell': f'{sheet}!{openpyxl.utils.get_column_letter(id_index+1)}{row_number}', 'message': 'Record_ID를 비어 있지 않은 문자열로 입력하세요.'})
                        continue
                    record = record.strip()
                    if record in seen:
                        self.errors.append({'cell': f'{sheet}!{openpyxl.utils.get_column_letter(id_index+1)}{row_number}', 'message': '같은 시트에서 Record_ID가 중복되었습니다.'})
                        continue
                    seen.add(record)
                    if supplement:
                        supplements.append((table, record, f'{sheet}!{openpyxl.utils.get_column_letter(id_index+1)}{row_number}'))
                    else:
                        source_records.add((table, record))
                    row = self.data.setdefault(table, {}).setdefault(record, {})
                    for col_index, column in enumerate(headers):
                        if column == 'Record_ID':
                            continue
                        cell = cells[col_index]
                        value = cell.value
                        location = f'{sheet}!{openpyxl.utils.get_column_letter(col_index+1)}{row_number}'
                        if column not in defs:
                            if value is not None:
                                self.unmapped.append({'cell': location, 'table': table, 'record': record, 'column': column or '(헤더 없음)', 'reason': '이 시험 프로필에서 사용하지 않는 컬럼'})
                            continue
                        key = table, record, column
                        self.locations[key] = location
                        if cell.data_type == 'f':
                            self.error(*key, '수식은 지원하지 않습니다. 값으로 붙여넣으세요.')
                            continue
                        if value is None or (isinstance(value, str) and not value.strip()):
                            continue
                        definition = defs[column]
                        try:
                            row[column] = parse_value(value, definition[1])
                            self.checks.append({'cell': location, 'table': table, 'record': record, 'column': column,
                                                'required': definition[2], 'mapping': definition[3]})
                        except (ValueError, TypeError, OverflowError):
                            self.error(*key, '입력 형식/범위 오류. 필요한 형식: '+definition[1])
            for table, record, location in supplements:
                if (table, record) not in source_records:
                    self.errors.append({'cell': location, 'message': '보완 행의 Record_ID와 연결되는 원본 테이블 행이 없습니다.'})
            for table, definitions in {**SOURCE, 'Files': FILES}.items():
                records = self.data.get(table, {})
                if not records and table != 'Files':
                    self.errors.append({'cell': table+'!A5', 'message': '최소 1개 데이터 행이 필요합니다.'})
                for record, row in records.items():
                    for column, _, required, _, _ in definitions + SUPPLEMENT.get(table, []):
                        if required and column not in row:
                            if (table, record, column) not in self.locations:
                                owner = next(s for s, (t, ds, _) in ROW_LAYOUT.items() if t == table and any(d[0] == column for d in ds))
                                self.locations[(table, record, column)] = owner+' (Record_ID='+record+')'
                            self.error(table, record, column, '필수값 또는 연결된 보완 행이 없습니다.')
            if '공구경로' in wb.sheetnames:
                ws = wb['공구경로']
                if [c.value for c in ws[4]][:8] != PATH_HEADERS:
                    self.errors.append({'cell': '공구경로!A4:H4', 'message': '경로 헤더가 변경되었습니다.'})
                else:
                    for row_number, cells in enumerate(ws.iter_rows(min_row=5, max_col=8), 5):
                        values = [c.value for c in cells]
                        if all(v is None for v in values):
                            continue
                        try:
                            if any(c.data_type == 'f' for c in cells):
                                raise ValueError()
                            rec = parse_value(values[0], 'text')
                            seq = parse_value(values[1], 'integer')
                            xyzijk = [parse_value(v, 'number') for v in values[2:]]
                            if not math.isclose(sum(x*x for x in xyzijk[3:]), 1, abs_tol=1e-6):
                                raise ValueError()
                            self.paths.setdefault(rec, []).append((seq, xyzijk, f'공구경로!A{row_number}'))
                        except (ValueError, TypeError, OverflowError):
                            self.errors.append({'cell': f'공구경로!A{row_number}:H{row_number}', 'message': '가공 Record_ID/순서/좌표/단위 공구축을 확인하세요.'})
        finally:
            wb.close()

    def index(self, table, columns):
        result = {}
        for record, row in self.data.get(table, {}).items():
            if not all(c in row for c in columns):
                continue
            key = tuple(row[c] for c in columns)
            if key in result:
                self.error(table, record, columns[-1], '업체 복합키가 중복되었습니다.')
            result[key] = record
        return result


def required_errors(obj, path='dt_asset'):
    errors = []
    if is_dataclass(obj):
        for field in fields(obj):
            value = getattr(obj, field.name)
            here = path + '.' + field.name
            if field.metadata.get('required') and (value is None or value == ''):
                errors.append(here)
            errors += required_errors(value, here)
    elif isinstance(obj, list):
        for i, child in enumerate(obj):
            errors += required_errors(child, f'{path}[{i}]')
    return errors


def stable_id(kind, *keys):
    digest = hashlib.sha256(json.dumps(keys, ensure_ascii=False).encode()).hexdigest()[:24]
    return kind + '-' + digest


def placement(name, row, prefix):
    return m.Axis2Placement3D(name=name,
        location=m.CartesianPoint(name=name+'-origin', coordinates=row[prefix+'_origin']),
        axis=m.Direction(name=name+'-axis', direction_ratios=row[prefix+'_axis']),
        ref_direction=m.Direction(name=name+'-ref', direction_ratios=row[prefix+'_ref_direction']))


def build(inp, workbook_path, part, rev, namespace, file_mode):
    data = inp.data
    indexes = {
        'Part_Info': inp.index('Part_Info', ['Part_No', 'Part_Rev']),
        'Operation_Info': inp.index('Operation_Info', ['Part_No', 'Part_Rev', 'Op_No']),
        'Tool_Milling': inp.index('Tool_Milling', ['공구번호', '공구사양']),
        'Machine_Milling': inp.index('Machine_Milling', ['설비코드', '설비번호']),
    }
    inp.index('Operation_Machining', ['Part_No', 'Part_Rev', 'Op_No', 'PGMNO', 'REVNO', 'SEQNO'])
    if inp.errors:
        return [], [], []
    parts = indexes['Part_Info']
    selected = (part, rev) if part is not None else (next(iter(parts)) if len(parts) == 1 else None)
    if selected not in parts:
        inp.errors.append({'cell': '업체입력', 'message': '파트/개정을 선택하세요. 여러 파트가 있으면 --part와 --rev가 필요합니다.'})
        return [], [], []
    p_rec = parts[selected]
    p = data['Part_Info'][p_rec]
    op_rows = {r: x for r, x in data['Operation_Info'].items() if (x['Part_No'], x['Part_Rev']) == selected}
    step_rows = {r: x for r, x in data['Operation_Machining'].items() if (x['Part_No'], x['Part_Rev']) == selected}
    if not op_rows or not step_rows:
        inp.errors.append({'cell': '업체입력', 'message': '선택한 파트에 공정과 상세 가공이 필요합니다.'})
    # Validate all joins, including records that will not be emitted for this selection.
    for r, x in data['Operation_Info'].items():
        for table, cols in [('Part_Info', ['Part_No', 'Part_Rev']), ('Machine_Milling', ['설비코드', '설비번호'])]:
            if tuple(x[c] for c in cols) not in indexes[table]:
                inp.error('Operation_Info', r, cols[-1], table + ' 연결 대상이 없습니다.')
    for r, x in data['Operation_Machining'].items():
        for table, cols in [('Operation_Info', ['Part_No', 'Part_Rev', 'Op_No']), ('Tool_Milling', ['공구번호', '공구사양'])]:
            if tuple(x[c] for c in cols) not in indexes[table]:
                inp.error('Operation_Machining', r, cols[-1], table + ' 연결 대상이 없습니다.')
        for source in ('spindle_source', 'feed_source'):
            if x[source] not in x:
                inp.error('Operation_Machining', r, x[source], '선택한 업체 컬럼의 값이 없습니다.')
        if x.get(x['spindle_source'], 1) == 0 or x.get(x['feed_source'], 1) <= 0:
            inp.error('Operation_Machining', r, 'feed_source', '절삭 이송은 양수, 주축은 0이 아닌 값이어야 합니다.')
        for prefix in ('feature', 'security', 'depth'):
            if abs(sum(a*b for a, b in zip(x[prefix+'_axis'], x[prefix+'_ref_direction']))) > 1e-6:
                inp.error('Operation_Machining', r, prefix+'_ref_direction', '평면의 Z축과 X축이 직교해야 합니다.')
    for r, x in data['Machine_Milling'].items():
        if x['축수'] != 5 or not (x['maximum_number_of_simultaneous_control_axes'] <= x['maximum_total_number_of_control_feed_axes'] == 5):
            inp.error('Machine_Milling', r, '축수', '5축 프로필의 총 이송축 수/동시 제어축 수를 확인하세요.')
    for table, rows, scope in [('Operation_Info', data['Operation_Info'], ['Part_No', 'Part_Rev']), ('Operation_Machining', data['Operation_Machining'], ['Part_No', 'Part_Rev', 'Op_No'])]:
        seen = set()
        for r, x in rows.items():
            key = tuple(x[c] for c in scope) + (x['execution_order'],)
            if key in seen:
                inp.error(table, r, 'execution_order', '같은 대상 안에서 실행 순서가 중복되었습니다.')
            seen.add(key)
    for r, points in inp.paths.items():
        if r not in data['Operation_Machining'] or len(points) < 2 or len({x[0] for x in points}) != len(points):
            inp.errors.append({'cell': points[0][2], 'message': '경로는 존재하는 가공 레코드에 연결하고, 중복 없는 순서로 2점 이상 입력하세요.'})
    if inp.errors:
        return [], [], []
    assets, links, staged = [], {}, []
    def asset(kind, keys, element):
        aid = stable_id(kind, *keys)
        element.element_id = aid
        obj = m.DtAsset(asset_global_id=namespace, id=aid, schema_version='v31',
                        asset_kind=m.DtAssetKind.INSTANCE, dt_elements=[element])
        assets.append(obj)
        return f'{namespace}/{aid}/{aid}'
    def ref(uri):
        return m.DtReference(element_id=stable_id('ref', uri), keys=[m.DtKeyValue(key=m.DtKeyType.DT_ELEMENT_FULLPATH, value=uri)])
    material = m.DtMaterial(standard_identifier=p['Spec_재질'], material_identifier=p['Part_Material'])
    mat_link = asset('material', [p['Spec_재질'], p['Part_Material']], material)
    wp = m.Workpiece(its_id=stable_id('workpiece', *selected), ref_dt_material=ref(mat_link))
    machine_records = {indexes['Machine_Milling'][(x['설비코드'], x['설비번호'])] for x in op_rows.values()}
    tool_records = {indexes['Tool_Milling'][(x['공구번호'], x['공구사양'])] for x in step_rows.values()}
    for r in sorted(machine_records):
        x = data['Machine_Milling'][r]
        nc_fields = [f.name for f in fields(m.NcController) if f.metadata.get('required')]
        nc = {k: x[k] for k in nc_fields if k not in ('controller_model', 'units')}
        machine = m.MachineToolSpecification(description=x['설비명 (모델)'], machine_class=m.MachineClass(x['machine_class']),
            device_id=m.DeviceId(id=x['설비코드']+':'+x['설비번호'], model_name=x['설비명 (모델)'], serial_number=x['serial_number'], manufacturer=x['메이커']),
            nc_controller_information=m.NcController(controller_model=x['컨트롤러'], units=m.UnitsType.METRIC, **nc))
        links[('Machine_Milling', r)] = asset('machine', [x['설비코드'], x['설비번호']], m.DtMachineTool(display_name=x['설비명 (모델)'], machine_tool=machine))
    for r in sorted(tool_records):
        x = data['Tool_Milling'][r]
        tool = m.DtCuttingTool13399(display_name=x['공구번호']+' '+x['공구사양'])
        for col, name, unit in [('날경 (mm)', 'effective_cutting_diameter', 'mm'), ('날수 (ea)', 'number_of_teeth', '1'), ('코너R', 'corner_radius', 'mm')]:
            if col in x:
                tool.numerical_value.append(m.NumericalValue(value_name=name, value_component=str(x[col]), unit_component=m.Unit(unit_name=unit)))
        if 'Maker' in x:
            tool.string_value.append(m.StringValue(value_name='manufacturer', value_specification=m.StringSelect(default_language_string=x['Maker'])))
        links[('Tool_Milling', r)] = asset('tool', [x['공구번호'], x['공구사양']], tool)
    main = m.Workplan(its_id=stable_id('main', *selected))
    project = m.DtProject(its_id=stable_id('project', *selected), display_name=p['Part_Name'], category='milling-registration-prototype',
        element_description='ELS offline prototype; not validated for machine execution', main_workplan=main, its_workpieces=[wp])
    project_uri = asset('project', selected, project)
    links[('Part_Info', p_rec)] = project_uri
    warnings = list(NOTES)
    for op_rec, op in sorted(op_rows.items(), key=lambda kv: kv[1]['execution_order']):
        plan_id = stable_id('workplan', *selected, op['Op_No']) + ':' + op['공정명']
        machine_rec = indexes['Machine_Milling'][(op['설비코드'], op['설비번호'])]
        plan = m.Workplan(its_id=plan_id, ref_dt_machine_tool=ref(links[('Machine_Milling', machine_rec)]))
        main.its_elements.append(plan)
        links[('Operation_Info', op_rec)] = project_uri
        steps = [(r, x) for r, x in step_rows.items() if x['Op_No'] == op['Op_No']]
        if not steps:
            inp.error('Operation_Info', op_rec, 'Op_No', '공정에 상세 가공이 없습니다.')
        for r, x in sorted(steps, key=lambda kv: kv[1]['execution_order']):
            sid = stable_id('step', *selected, x['Op_No'], x['PGMNO'], x['REVNO'], x['SEQNO'])
            tr = indexes['Tool_Milling'][(x['공구번호'], x['공구사양'])]
            tx = data['Tool_Milling'][tr]
            tool = m.MillingCuttingTool(its_id=x['TOOLNO'], overall_assembly_length=x['overall_assembly_length'],
                maximum_depth_of_cut=x['maximum_depth_of_cut'], effective_cutting_diameter=tx['날경 (mm)'], number_of_effective_teeth=tx['날수 (ea)'], edge_radius=tx.get('코너R'))
            tech = m.MillingTechnology(feedrate=x[x['feed_source']] / (60000 if x['feed_unit'] == 'mm/min' else 1),
                spindle=x[x['spindle_source']] / (60 if x['spindle_unit'] == 'rpm' else 1), feedrate_reference=m.ToolReferencePoint(x['feedrate_reference']),
                **{k: x[k] for k in ('synchronize_spindle_with_feed', 'inhibit_feedrate_override', 'inhibit_spindle_override')})
            operation = m.FreeformOperation(its_id=sid+'-operation', its_tool=tool, its_technology=tech,
                its_machine_functions=m.MillingMachineFunctions(**{k:x[k] for k in ('coolant', 'through_spindle_coolant', 'chip_removal')}),
                ref_dt_cutting_tool=ref(links[('Tool_Milling', tr)]))
            if r in inp.paths:
                points = sorted(inp.paths[r])
                operation.its_toolpath = m.ToolpathList(its_list=[m.CutterLocationTrajectory(its_priority=True,
                    its_type=m.ToolpathType.TRAJECTORY_PATH,
                    basiccurve=m.Polyline(name=sid+'-path', points=[m.CartesianPoint(name=f'p{i}', coordinates=v[:3]) for i,v,_ in points]),
                    its_toolaxis=m.Polyline(name=sid+'-axis', points=[m.CartesianPoint(name=f'a{i}', coordinates=v[3:]) for i,v,_ in points]))])
            else:
                warnings.append(f'{r}: 경로·공구축 미입력. 가공 실행용 데이터는 미완성입니다.')
            feature = m.ToolpathFeature(its_id=sid+'-feature', its_workpiece=wp, its_operations=[operation], feature_placement=placement(sid+'-feature', x, 'feature'),
                depth=m.Plane(name=sid+'-depth', position=placement(sid+'-depth', x, 'depth')))
            plan.its_elements.append(m.MachiningWorkingstep(its_id=sid, its_feature=feature, its_operation=operation,
                its_secplane=m.Plane(name=sid+'-security', position=placement(sid+'-security', x, 'security'))))
    # Only files attached to the selected project or its actually used resources are emitted.
    for r, x in data.get('Files', {}).items():
        targets = {'Part_File': ('Part_Info', 'Operation_Info'), 'Tool_File': ('Tool_Milling',), 'Machine_File': ('Machine_Milling',)}[x['source_table']]
        matches = [(t, x['target_record']) for t in targets if x['target_record'] in data.get(t, {})]
        if len(matches) != 1:
            inp.error('Files', r, 'target_record', '파일 연결 대상이 없거나 모호합니다.')
            continue
        if matches[0] not in links:
            continue
        path = Path(x['File_Path']).expanduser()
        if not path.is_absolute():
            path = workbook_path.parent / path
        if not path.is_file():
            inp.error('Files', r, 'File_Path', '로컬 파일을 읽을 수 없습니다.')
            continue
        if '/' in x['File_Name'] or '\\' in x['File_Name'] or x['File_Name'] in ('.', '..'):
            inp.error('Files', r, 'File_Name', '파일명에 경로를 넣지 마세요.')
            continue
        if not re.fullmatch(r'[\w.+-]+/[\w.+-]+', x['File_Type']):
            inp.error('Files', r, 'File_Type', 'MIME 형식이 필요합니다. 예: text/plain')
            continue
        digest = hashlib.sha256()
        with path.open('rb') as stream:
            for chunk in iter(lambda: stream.read(1024*1024), b''):
                digest.update(chunk)
        sha = digest.hexdigest()
        value = x.get('gridfs_id', '')
        if file_mode == 'gridfs' and not re.fullmatch('[0-9a-fA-F]{24}', value):
            inp.error('Files', r, 'gridfs_id', '실제 업로드 결과의 24자리 ObjectId가 필요합니다.')
            continue
        if file_mode == 'local':
            value = 'pending-upload:' + sha
        packed = 'files/' + sha + Path(x['File_Name']).suffix
        reference = ref(links[matches[0]])
        if matches[0][0] == 'Operation_Info':
            target_op = data['Operation_Info'][matches[0][1]]
            reference.keys.append(m.DtKeyValue(key=m.DtKeyType.WORKPLAN, value=stable_id('workplan', *selected, target_op['Op_No'])+':'+target_op['공정명']))
        f = m.DtFile(display_name=x['File_Name'], content_type=x['File_Type'], category=x['Doc_Type'], value=value, path=packed,
                     reference=[reference], properties=[m.KeyValue(key='sha256', value=sha)])
        if x.get('Doc_Rev'):
            f.properties.append(m.KeyValue(key='Doc_Rev', value=x['Doc_Rev']))
        uri = asset('file', [links[matches[0]], x['File_Name'], x.get('Doc_Rev', ''), sha], f)
        staged.append({'source': path, 'path': packed, 'sha256': sha, 'asset_uri': uri, 'upload_required': file_mode == 'local'})
    if file_mode == 'local' and staged:
        warnings.append('dt_file.value는 pending-upload 임시값입니다. 서버 API 등록 전에 파일 업로드와 ObjectId 치환이 필요합니다.')
    if file_mode == 'gridfs' and staged:
        warnings.append('GridFS ObjectId 형식만 확인했습니다. 서버 파일 존재 여부와 파일 내용 일치는 확인하지 않았습니다.')
    if not staged:
        warnings.append('선택 대상에 연결된 파일이 없어 dt_file은 생성하지 않았습니다.')
    for obj in assets:
        for missing in required_errors(obj):
            inp.errors.append({'cell': '생성 모델', 'message': '필수 필드 누락: '+missing})
    if len({a.id for a in assets}) != len(assets):
        inp.errors.append({'cell': '파일입력', 'message': '같은 대상·파일명·개정·내용의 파일이 중복되었습니다.'})
    return assets, staged, warnings


def run(input_path, output, part=None, rev=None, namespace='urn:els-local', file_mode='local'):
    input_path, output = Path(input_path).resolve(), Path(output).resolve()
    if output.exists():
        raise InputError('출력 폴더가 이미 있습니다. 새 폴더를 지정하세요.')
    if bool(part is None) != bool(rev is None):
        raise InputError('--part와 --rev를 함께 지정하세요.')
    if file_mode not in ('local', 'gridfs') or not re.fullmatch(r'[A-Za-z][A-Za-z0-9+.-]*:[^\s<>"\x00-\x1f]+', namespace):
        raise InputError('file_mode 또는 namespace URI가 잘못되었습니다.')
    inp = Input()
    inp.read(input_path)
    assets, files_to_copy, warnings = ([], [], []) if inp.errors else build(inp, input_path, part, rev, namespace.rstrip('/'), file_mode)
    xmls = {}
    if not inp.errors:
        serializer = XmlSerializer(config=SerializerConfig(indent='  '))
        parser = XmlParser()
        for obj in assets:
            xml = serializer.render(obj, ns_map={None: m.__NAMESPACE__})
            parsed = parser.parse(io.BytesIO(xml.encode('utf-8')), m.DtAsset)
            if parsed != obj or required_errors(parsed):
                raise InputError('XML 왕복 검증 실패. 파일을 생성하지 않았습니다.')
            xmls[obj.id + '.xml'] = xml
    # Explain which validated inputs were actually selected for this output.
    selected_part = (part, rev)
    if part is None and len(inp.data.get('Part_Info', {})) == 1:
        p = next(iter(inp.data['Part_Info'].values()))
        selected_part = p.get('Part_No'), p.get('Part_Rev')
    for check in inp.checks:
        row = inp.data[check['table']][check['record']]
        if 'Part_No' in row and (row.get('Part_No'), row.get('Part_Rev')) != selected_part:
            check['output_use'] = '선택하지 않은 파트'
        elif check['table'] == 'Operation_Machining' and (check['column'].startswith('SPINDLE_') or check['column'].startswith('FEED_')) and check['column'] not in (row.get('spindle_source'), row.get('feed_source')):
            check['output_use'] = '주축/이송 후보 중 선택하지 않은 컬럼'
        else:
            check['output_use'] = '값/연결/변환 규칙에 사용. 마스터·파일은 선택 파트에서 참조될 때만 출력'
    report = {'ok': not inp.errors, 'machine_ready': False, 'profile': '5-axis-milling-freeform-registration-prototype',
              'file_mode': file_mode, 'server_file_existence_verified': False,
              'errors': inp.errors, 'warnings': warnings, 'input_checks': inp.checks,
              'unmapped_columns': inp.unmapped, 'xml_files': list(xmls),
              'model_required_fields_verified': not inp.errors,
              'assets': [{'id':a.id, 'type':type(a.dt_elements[0]).__name__} for a in assets] if not inp.errors else []}
    output.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix='.els-convert-', dir=output.parent))
    try:
        (staging / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        if not inp.errors:
            for name, xml in xmls.items():
                (staging / name).write_text(xml, encoding='utf-8')
            for item in files_to_copy:
                dest = staging / item['path']
                dest.parent.mkdir(exist_ok=True)
                shutil.copyfile(item['source'], dest)
                with dest.open('rb') as stream:
                    if hashlib.file_digest(stream, 'sha256').hexdigest() != item['sha256']:
                        raise InputError('변환 중 원본 파일이 변경되었습니다. 다시 실행하세요.')
            (staging / 'file_manifest.json').write_text(json.dumps([{k:v for k,v in item.items() if k != 'source'} for item in files_to_copy], ensure_ascii=False, indent=2), encoding='utf-8')
        staging.rename(output)
    finally:
        if staging.exists():
            shutil.rmtree(staging)
    return report


def main():
    parser = argparse.ArgumentParser(description='로컬 Excel → DT XML. 외부 전송/DB 변경 없음.')
    parser.add_argument('input', type=Path)
    parser.add_argument('--out', type=Path, required=True, help='새 출력 폴더')
    parser.add_argument('--part', help='Part_No (문자열)')
    parser.add_argument('--rev', help='Part_Rev (문자열)')
    parser.add_argument('--namespace', default='urn:els-local', help='자산 식별 범위 URI. 반복 변환 시 같은 값 사용')
    parser.add_argument('--file-mode', choices=['local', 'gridfs'], default='local')
    args = parser.parse_args()
    try:
        report = run(args.input, args.out, args.part, args.rev, args.namespace, args.file_mode)
    except (InputError, OSError, ValueError, KeyError, BadZipFile, openpyxl.utils.exceptions.InvalidFileException) as exc:
        # Input records/paths may be confidential. Do not print arbitrary exception bodies.
        print(str(exc) if isinstance(exc, InputError) else '입력 파일/형식을 읽지 못했습니다. Excel 양식과 로컬 파일 접근 권한을 확인하세요.', file=sys.stderr)
        return 2
    print(f"XML {len(report['xml_files'])}개, 오류 {len(report['errors'])}개. 결과: {args.out / 'report.json'}")
    for error in report['errors'][:15]:
        print(f"- {error['cell']}: {error.get('column', '')} {error['message']}")
    return 0 if report['ok'] else 2


if __name__ == '__main__':
    raise SystemExit(main())

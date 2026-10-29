"""Exact ELS columns/order at row 1, with DT supplements in separate sheets."""
import json
import math
from pathlib import Path

import openpyxl
from schema import SOURCE, SUPPLEMENT, FILES

SPEC = json.loads(Path(__file__).with_name('original_columns.json').read_text(encoding='utf-8'))
KEYS = {
    'Part_Info': ['Part_No', 'Part_Rev'],
    'Operation_Info': ['Part_No', 'Part_Rev', 'Op_No'],
    'Operation_Machining': ['Operation_Machining_SN'],
    'Tool_Milling': ['공구번호', '공구사양'],
    'Machine_Milling': ['설비코드', '설비번호'],
    'Part_File': ['Part_File_SN'], 'Tool_File': ['Tool_File_SN'], 'Machine_File': ['Machine_File_SN'],
}
DT_SHEETS = {
    'DT_공정': ('Operation_Info', SUPPLEMENT['Operation_Info']),
    'DT_가공': ('Operation_Machining', SUPPLEMENT['Operation_Machining']),
    'DT_장비': ('Machine_Milling', SUPPLEMENT['Machine_Milling']),
}
FILE_HEADERS = ['source_table', 'File_SN', 'content_type', 'local_path', 'gridfs_id']
PATH_HEADERS = ['Operation_Machining_SN', '순서', 'X_mm', 'Y_mm', 'Z_mm', 'I', 'J', 'K']


def serial(value):
    if isinstance(value, bool) or value is None:
        raise ValueError('SN이 없습니다.')
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    if isinstance(value, str) and value.strip():
        return value.strip()
    raise ValueError('SN은 정수 또는 문자열이어야 합니다.')


def record_key(table, row):
    values = [serial(row[k]) if k.endswith('_SN') else row[k] for k in KEYS[table]]
    if not all(isinstance(v, str) and v.strip() for v in values):
        raise ValueError('복합키는 비어 있지 않은 문자열이어야 합니다.')
    return table + ':' + json.dumps(values, ensure_ascii=False, separators=(',', ':'))


def read_original(inp, filename, parse):
    wb = openpyxl.load_workbook(filename, data_only=False, read_only=True)
    locations = {}
    file_rows = []
    dt_file_rows = {}

    def rows(sheet, expected, optional=False):
        if sheet not in wb.sheetnames:
            if not optional:
                inp.errors.append({'cell': sheet, 'message': '시트가 없습니다.'})
            return []
        ws = wb[sheet]
        actual = [c.value for c in ws[1]]
        while actual and actual[-1] is None:
            actual.pop()
        if actual != expected:
            inp.errors.append({'cell': sheet+'!1', 'message': '컬럼명 또는 순서가 원본 양식과 다릅니다.'})
            return []
        result = []
        for n, cells in enumerate(ws.iter_rows(min_row=2, max_col=len(expected)), 2):
            if all(c.value is None or c.value == '' for c in cells):
                continue
            item = {}
            for i, (name, cell) in enumerate(zip(expected, cells), 1):
                loc = f'{sheet}!{openpyxl.utils.get_column_letter(i)}{n}'
                locations[(sheet, n, name)] = loc
                if cell.data_type == 'f':
                    inp.errors.append({'cell': loc, 'message': '수식 대신 값으로 붙여넣으세요.'})
                else:
                    item[name] = cell.value
            result.append((n, item))
        return result

    def identifier(table, row, sheet, n):
        try:
            # Key normalisation is limited to SerialNo columns; text IDs keep zeros.
            return record_key(table, row)
        except (KeyError, ValueError):
            inp.errors.append({'cell': sheet+f'!{n}', 'message': '원본 기본키가 누락되었거나 형식이 다릅니다: '+', '.join(KEYS[table])})
            return None

    def put(table, rec, column, value, definition, loc):
        inp.locations[(table, rec, column)] = loc
        if value is None or (isinstance(value, str) and not value.strip()):
            return
        try:
            inp.data[table][rec][column] = parse(value, definition[1])
            inp.checks.append({'cell': loc, 'table': table, 'record': rec, 'column': column,
                                'required': definition[2], 'mapping': definition[3]})
        except (ValueError, TypeError, OverflowError):
            inp.error(table, rec, column, '입력 형식/범위 오류: '+definition[1])

    try:
        for table, spec in SPEC['tables'].items():
            expected = [c['name'] for c in spec['columns']]
            seen = set()
            for n, row in rows(table, expected, optional=table not in SOURCE):
                if table not in KEYS:
                    # Optional reference tables are preserved, never silently mapped.
                    for column, value in row.items():
                        if value is not None:
                            inp.unmapped.append({'cell': locations[(table,n,column)], 'table':table,'column':column,'reason':'현재 변환 프로필에서 참고용으로만 보관'})
                    continue
                rec = identifier(table, row, table, n)
                if rec is None:
                    continue
                if rec in seen:
                    inp.errors.append({'cell': table+f'!{n}', 'message': '원본 기본키가 중복되었습니다.'})
                    continue
                seen.add(rec)
                if table.endswith('_File'):
                    file_rows.append((table, rec, n, row))
                    continue
                inp.data.setdefault(table, {})[rec] = {}
                defs = {d[0]: d for d in SOURCE[table]}
                for column, value in row.items():
                    if column in defs:
                        put(table, rec, column, value, defs[column], locations[(table,n,column)])
                    elif value is not None and column not in KEYS[table]:
                        inp.unmapped.append({'cell': locations[(table,n,column)], 'table':table,'record':rec,'column':column,'reason':'원본 컬럼을 유지하며 DT 변환에는 사용하지 않음'})
        for sheet, (table, defs) in DT_SHEETS.items():
            expected = KEYS[table] + [d[0] for d in defs]
            seen = set()
            for n, row in rows(sheet, expected, optional=True):
                rec = identifier(table, row, sheet, n)
                if rec is None:
                    continue
                if rec in seen:
                    inp.errors.append({'cell':sheet+f'!{n}','message':'DT 보완 기본키가 중복되었습니다.'})
                    continue
                seen.add(rec)
                if rec not in inp.data.get(table, {}):
                    inp.errors.append({'cell':sheet+f'!{n}','message':'원본 테이블의 연결 대상이 없습니다.'})
                    continue
                for d in defs:
                    put(table,rec,d[0],row.get(d[0]),d,locations[(sheet,n,d[0])])
        for n, row in rows('DT_파일', FILE_HEADERS, optional=True):
            try:
                table = row['source_table']
                if table not in ('Part_File','Tool_File','Machine_File'):
                    raise ValueError()
                rec = record_key(table, {KEYS[table][0]:row['File_SN']})
                if rec in dt_file_rows:
                    raise ValueError()
                dt_file_rows[rec] = (n,row)
            except (ValueError,KeyError):
                inp.errors.append({'cell':f'DT_파일!{n}','message':'파일 테이블·원본 SN을 확인하세요. 동일 키는 중복 불가합니다.'})
        file_keys = {r for _,r,_,_ in file_rows}
        for rec, (n,_) in dt_file_rows.items():
            if rec not in file_keys:
                inp.errors.append({'cell':f'DT_파일!{n}','message':'원본 파일 테이블에 대응하는 SN이 없습니다.'})
        for table,rec,n,row in file_rows:
            inp.data.setdefault('Files', {})[rec] = {}
            try:
                if table == 'Part_File':
                    target = 'Operation_Info' if row.get('Op_No') not in (None,'') else 'Part_Info'
                    target_rec = record_key(target,row)
                elif table == 'Tool_File':
                    target_rec = record_key('Tool_Milling',row)
                else:
                    target_rec = record_key('Machine_Milling', {'설비코드':row.get('설비코드'), '설비번호':row.get('번호')})
            except (ValueError,KeyError):
                inp.errors.append({'cell':f'{table}!{n}','message':'파일 연결용 원본 키를 확인하세요.'})
                continue
            dt_n, extra = dt_file_rows.get(rec, (None, {}))
            merged = dict(row, source_table=table, target_record=target_rec,
                          File_Type=extra.get('content_type'), gridfs_id=extra.get('gridfs_id'))
            if extra.get('local_path'):
                merged['File_Path'] = extra['local_path']
            for d in FILES:
                column=d[0]
                loc=locations.get((table,n,column),f'{table}!{n}')
                if column == 'File_Type':
                    loc=f'DT_파일!C{dt_n}' if dt_n else f'DT_파일 ({KEYS[table][0]}={row[KEYS[table][0]]})'
                elif column=='gridfs_id' and dt_n:
                    loc=f'DT_파일!E{dt_n}'
                elif column=='File_Path' and extra.get('local_path'):
                    loc=f'DT_파일!D{dt_n}'
                put('Files',rec,column,merged.get(column),d,loc)
        for n,row in rows('DT_공구경로',PATH_HEADERS,optional=True):
            try:
                rec=record_key('Operation_Machining',row)
                seq=parse(row['순서'],'integer')
                xyzijk=[parse(row[k],'number') for k in PATH_HEADERS[2:]]
                if not math.isclose(sum(v*v for v in xyzijk[3:]),1,abs_tol=1e-6):
                    raise ValueError()
                inp.paths.setdefault(rec,[]).append((seq,xyzijk,f'DT_공구경로!A{n}'))
            except (ValueError,KeyError,TypeError,OverflowError):
                inp.errors.append({'cell':f'DT_공구경로!{n}','message':'원본 가공 SN, 순서, 좌표와 단위 공구축을 확인하세요.'})
        for table, defs in {**SOURCE,'Files':FILES}.items():
            records=inp.data.get(table,{})
            if not records and table!='Files':
                inp.errors.append({'cell':table+'!A2','message':'최소 1개 원본 데이터 행이 필요합니다.'})
            for rec,row in records.items():
                for d in defs+SUPPLEMENT.get(table,[]):
                    if d[2] and d[0] not in row:
                        if (table,rec,d[0]) not in inp.locations:
                            owner=next((s for s,(t,ds) in DT_SHEETS.items() if t==table and any(f[0]==d[0] for f in ds)),table)
                            inp.locations[(table,rec,d[0])]=owner+' (원본 키: '+rec+')'
                        inp.error(table,rec,d[0],'필수값 또는 DT 보완 행이 없습니다.')
    finally:
        wb.close()

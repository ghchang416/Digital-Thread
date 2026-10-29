"""Database-style workbook integration tests, including one-to-many tools."""
import copy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import convert as c

ROOT = Path(__file__).resolve().parents[3]
EXAMPLE = ROOT / 'outputs/els_dt_converter/가상예제_DB행입력.xlsx'
BLANK = ROOT / 'outputs/els_dt_converter/업체입력_빈양식_DB행입력.xlsx'


class RowWorkbookTests(unittest.TestCase):
    def read(self, wb=None):
        obj = c.Input()
        if wb is None:
            obj.read(EXAMPLE)
        else:
            with patch.object(c.openpyxl, 'load_workbook', return_value=wb):
                obj.read(EXAMPLE)
        return obj

    def build(self, inp):
        return c.build(inp, EXAMPLE, None, None, 'urn:test', 'local')

    def test_multiple_operations_tools_and_reuse_roundtrip(self):
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp)/'out'
            report = c.run(EXAMPLE, out)
            self.assertTrue(report['ok'], report['errors'])
            self.assertEqual(len(report['assets']), 7)
            self.assertEqual(sum(a['type']=='DtCuttingTool13399' for a in report['assets']), 3)
            project_path = next(out.glob('project-*.xml'))
            project = c.XmlParser().parse(project_path, c.m.DtAsset).dt_elements[0]
            plans = project.main_workplan.its_elements
            self.assertEqual([len(p.its_elements) for p in plans], [4,2])
            refs = [[s.its_operation.ref_dt_cutting_tool.keys[0].value for s in p.its_elements] for p in plans]
            self.assertEqual(refs[0][0], refs[0][3])
            self.assertEqual(refs[0][1:], [refs[1][0], refs[1][1], refs[0][0]])
            self.assertEqual(len(set(refs[0])), 3)
            self.assertEqual(c.required_errors(project), [])

    def test_arbitrary_number_of_steps_without_tool_columns(self):
        inp = self.read()
        self.assertFalse(inp.errors)
        base = inp.data['Operation_Machining']['WS1']
        for i in range(7,127):
            inp.data['Operation_Machining'][f'WS{i}'] = dict(base, execution_order=i, SEQNO=i*10)
        assets, _, _ = self.build(inp)
        self.assertFalse(inp.errors)
        project = next(a.dt_elements[0] for a in assets if isinstance(a.dt_elements[0], c.m.DtProject))
        self.assertEqual(sum(len(p.its_elements) for p in project.main_workplan.its_elements), 126)
        self.assertEqual(sum(isinstance(a.dt_elements[0], c.m.DtCuttingTool13399) for a in assets), 3)

    def test_missing_supplement_row(self):
        wb = c.openpyxl.load_workbook(EXAMPLE)
        wb['Machining_Conditions'].delete_rows(6)
        inp = self.read(wb)
        self.assertTrue(any(e.get('record')=='WS2' and '필수값' in e['message'] for e in inp.errors))

    def test_orphan_supplement_row(self):
        wb = c.openpyxl.load_workbook(EXAMPLE)
        wb['Machining_Geometry']['A5'] = 'MISSING'
        inp = self.read(wb)
        self.assertTrue(any(e['cell']=='Machining_Geometry!A5' and '원본' in e['message'] for e in inp.errors))

    def test_duplicate_record_id(self):
        wb = c.openpyxl.load_workbook(EXAMPLE)
        wb['Operation_Machining']['A6'] = 'WS1'
        inp = self.read(wb)
        self.assertTrue(any(e['cell']=='Operation_Machining!A6' and '중복' in e['message'] for e in inp.errors))

    def test_formula_and_missing_value_have_exact_cells(self):
        wb = c.openpyxl.load_workbook(EXAMPLE)
        wb['Part_Info']['B5'] = '=1+1'
        wb['Part_Info']['E5'] = None
        inp = self.read(wb)
        self.assertTrue(any(e['cell']=='Part_Info!B5' and '수식' in e['message'] for e in inp.errors))
        self.assertTrue(any(e['cell']=='Part_Info!E5' and '필수값' in e['message'] for e in inp.errors))

    def test_blank_has_no_phantom_records(self):
        with tempfile.TemporaryDirectory() as temp:
            out=Path(temp)/'out'
            report=c.run(BLANK,out)
            self.assertFalse(report['ok'])
            self.assertEqual(list(out.glob('*.xml')), [])
            self.assertTrue(any(e['cell']=='Part_Info!A5' for e in report['errors']))

    def test_physical_row_order_does_not_control_execution(self):
        inp = self.read()
        inp.data['Operation_Machining'] = dict(reversed(list(inp.data['Operation_Machining'].items())))
        assets, _, _ = self.build(inp)
        project = next(a.dt_elements[0] for a in assets if isinstance(a.dt_elements[0], c.m.DtProject))
        plan=project.main_workplan.its_elements[0]
        self.assertEqual([s.its_operation.its_tool.its_id for s in plan.its_elements], ['T01','T02','T03','T01'])

    def test_preserve_business_keys(self):
        inp = self.read()
        self.assertEqual(inp.data['Part_Info']['P1']['Part_Rev'], '01')
        self.assertEqual(inp.data['Operation_Info']['OP1']['Op_No'], '010')
        self.assertEqual(inp.data['Machine_Milling']['M1']['설비번호'], '001')

    def test_append_after_formatted_range_is_read(self):
        wb=c.openpyxl.load_workbook(EXAMPLE)
        for name in ('Operation_Machining','Machining_Conditions','Machining_Geometry'):
            ws=wb[name]
            values=[cell.value for cell in ws[5]]
            values[0]='WS7'
            if name=='Operation_Machining':
                headers=[cell.value for cell in ws[4]]
                values[headers.index('SEQNO')]=70
                values[headers.index('execution_order')]=7
            for i,v in enumerate(values,1):
                ws.cell(150,i,v)
        inp=self.read(wb)
        self.assertFalse(inp.errors)
        self.assertEqual(len(inp.data['Operation_Machining']),7)
        assets,_,_=self.build(inp)
        self.assertFalse(inp.errors)
        project=next(a.dt_elements[0] for a in assets if isinstance(a.dt_elements[0],c.m.DtProject))
        self.assertEqual([len(p.its_elements) for p in project.main_workplan.its_elements],[5,2])

    def test_sparse_empty_cells_still_report_coordinates(self):
        from openpyxl.cell.read_only import EMPTY_CELL
        wb=c.openpyxl.load_workbook(EXAMPLE)
        ws=wb['Part_Info']
        ws['D5']=None
        original=ws.iter_rows
        def sparse(*args,**kwargs):
            for row in original(*args,**kwargs):
                yield tuple(EMPTY_CELL if cell.value is None else cell for cell in row)
        with patch.object(ws,'iter_rows',side_effect=sparse):
            inp=self.read(wb)
        self.assertTrue(any(e['cell']=='Part_Info!D5' for e in inp.errors))


if __name__ == '__main__':
    unittest.main()
